#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_lint_fixtures.py — Task 9.2 комплекта report-style-kit. Фикстуры lint_text.

13 кейсов:
  8 негативных одно-блочных манифестов (по одной категории блэклиста):
    first_person, agent_phrases, marketing, office_anglicisms, emoji,
    markdown_residue, exclamations, placeholders
    → каждый должен дать lint_text exit 1 с этой категорией в отчёте;
  5 позитивных (тексты эталона: аннотация [5], тезис [7], результат с числами
    [51], verdict-буллет [103], lead_in «Бизнес-вопрос:» [8] + текст [9])
    → каждый должен дать exit 0 (предупреждения допустимы).

Профиль: tests/work_extract_check/style-profile.json. Эталон читается только.
Итог: «FIXTURES PASS 13/13», exit 0; иначе список провалившихся кейсов, exit 1.

Запуск: python3 tests/test_lint_fixtures.py
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

import docx

KIT = Path(__file__).resolve().parents[1]
SAMPLE = KIT.parent / "CR_Predictor_Report.docx"
PROFILE = "tests/work_extract_check/style-profile.json"

META = {"title": "Проверка фикстур линтера", "subtitle": "test_lint_fixtures",
        "project": "fixture", "version": "V0", "period": "2026-09"}

# (категория, текст-нарушение) — ровно одна категория на кейс
NEGATIVE = [
    ("first_person", "Мы применили модель к данным за квартал."),
    ("agent_phrases", "Давайте рассмотрим структуру отчёта."),
    ("marketing", "Это уникальное решение для всей отрасли."),
    ("office_anglicisms", "Получили фидбек от технолога цеха."),
    ("emoji", "Рост 📊 показан на диаграмме."),
    ("markdown_residue", "**Важно** учитывать контекст измерений."),
    ("exclamations", "Результат отличный!"),
    ("placeholders", "Значение показателя не заполнено: $unresolved"),
]


def manifest_of(blocks: list) -> dict:
    return {"meta": META,
            "sections": [{"heading": "1. Фикстура", "level": 1, "blocks": blocks}]}


def reference_texts() -> dict:
    d = docx.Document(str(SAMPLE))
    return {i: d.paragraphs[i].text for i in (5, 7, 8, 9, 51, 103)}


def run_lint(manifest: dict, workdir: Path, name: str) -> tuple[int, str]:
    m_path = workdir / f"{name}.json"
    out_path = workdir / f"{name}_report.txt"
    m_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2),
                      encoding="utf-8")
    proc = subprocess.run(
        [sys.executable, "scripts/lint_text.py", "--manifest", str(m_path),
         "--profile", PROFILE, "--out", str(out_path)],
        cwd=KIT, capture_output=True, text=True)
    report = out_path.read_text(encoding="utf-8") if out_path.is_file() \
        else proc.stdout
    return proc.returncode, report


def main() -> int:
    if not SAMPLE.is_file():
        print(f"[fixtures] эталон не найден: {SAMPLE}", file=sys.stderr)
        return 1
    failures: list = []
    total = 0
    texts = reference_texts()

    with tempfile.TemporaryDirectory(prefix="lint_fixtures_") as tmp:
        workdir = Path(tmp)

        # --- негативные: exit 1 + категория в отчёте ------------------------
        for cat, text in NEGATIVE:
            total += 1
            code, report = run_lint(manifest_of(
                [{"type": "thesis", "text": text}]), workdir, f"neg_{cat}")
            if code != 1:
                failures.append(f"{cat}: ожидался exit 1, получен {code}")
            elif cat not in report:
                failures.append(f"{cat}: категории нет в отчёте линтера")
            else:
                print(f"  ok  {cat}: exit 1, категория в отчёте")

        # --- позитивные: exit 0 ----------------------------------------------
        positive = [
            ("thesis_anno", [{"type": "thesis", "text": texts[5]}]),
            ("thesis_ctx", [{"type": "thesis", "text": texts[7]}]),
            ("thesis_result", [{"type": "thesis", "text": texts[51]}]),
            ("verdict", [{"type": "verdict", "text": texts[103]}]),
            ("lead_in", [{"type": "lead_in", "label": texts[8],
                          "text": texts[9]}]),
        ]
        for name, blocks in positive:
            total += 1
            code, report = run_lint(manifest_of(blocks), workdir, f"pos_{name}")
            if code != 0:
                failures.append(f"{name}: ожидался exit 0, получен {code}\n{report}")
            else:
                print(f"  ok  {name}: exit 0")

    if failures:
        print(f"FIXTURES FAIL: {len(failures)}/{total}")
        for f in failures:
            print("  -", f)
        return 1
    print(f"FIXTURES PASS {total}/{total}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
